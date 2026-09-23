from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPORT_LIMIT = 500


def _set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color="D9D9D9", size="6"):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def _repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _keep_row_together(row):
    properties = row._tr.get_or_add_trPr()
    cannot_split = OxmlElement("w:cantSplit")
    cannot_split.set(qn("w:val"), "true")
    properties.append(cannot_split)


def _set_run_font(run, size, *, bold=False, color="000000"):
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")


def _remove_paragraph_borders(paragraph_or_style):
    element = paragraph_or_style._element
    properties = element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _write_cell(cell, value, *, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    _set_run_font(paragraph.add_run(str(value or "-")), 8.5, bold=bold, color=color)
    _set_cell_margins(cell)


def build_reservation_report(rows, *, section_title, period_text, generated_at, total_count, truncated=False):
    """Build a compact, printable Uzbek Latin DOCX reservation report."""
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    document.core_properties.title = "Band qilish arizalari hisoboti"
    document.core_properties.subject = section_title

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    _remove_paragraph_borders(title_style)
    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    _remove_paragraph_borders(title)
    _set_run_font(title.add_run("Band qilish arizalari hisoboti"), 18, bold=True)

    summary = document.add_paragraph()
    summary.paragraph_format.space_after = Pt(8)
    summary.paragraph_format.keep_with_next = True
    _set_run_font(summary.add_run(f"Bo‘lim: {section_title}\n"), 10, bold=True)
    _set_run_font(summary.add_run(f"Davr: {period_text}\n"), 9.5)
    _set_run_font(summary.add_run(f"Jami arizalar: {total_count} · Fayldagi yozuvlar: {len(rows)}\n"), 9.5)
    _set_run_font(summary.add_run(f"Hisobot yaratilgan: {generated_at}"), 9.5)
    if truncated:
        _set_run_font(summary.add_run("\nEslatma: fayl tez ochilishi uchun eng yangi 500 ta ariza kiritildi."), 9, bold=True)

    headers = ("T/r", "Ariza", "Mijoz", "Tashrif", "Joy", "Holat", "Izoh va sabab")
    widths = (0.35, 0.75, 1.25, 0.9, 1.05, 1.0, 1.55)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    _set_table_borders(table)

    header = table.rows[0]
    _repeat_table_header(header)
    _keep_row_together(header)
    for index, (cell, label, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        _set_cell_fill(cell, "4B2A1E")
        alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 1, 3, 5) else WD_ALIGN_PARAGRAPH.LEFT
        _write_cell(cell, label, bold=True, color="FFFFFF", align=alignment)

    for row_index, item in enumerate(rows, start=1):
        row = table.add_row()
        _keep_row_together(row)
        values = (
            str(row_index),
            f"{item['public_number']}\n{item['created_at']}",
            f"{item['customer_name']}\n{item['phone']}",
            f"{item['visit_at']}\n{item['guests_count']} kishi",
            f"{item['space']}\n{item['occasion']}",
            f"{item['status']}\n{item['manager']}",
            item["notes"],
        )
        for cell_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            if row_index % 2 == 0:
                _set_cell_fill(cell, "F7F3EF")
            alignment = WD_ALIGN_PARAGRAPH.CENTER if cell_index in (0, 1, 3, 5) else WD_ALIGN_PARAGRAPH.LEFT
            _write_cell(cell, value, align=alignment)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
