import uuid
from datetime import date, datetime, time, timedelta
from io import BytesIO
from unittest.mock import patch

from asgiref.sync import async_to_sync
from docx import Document
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.test import override_settings
from django.urls import reverse
from django.utils import timezone

from menu.models import RestaurantSettings

from .models import Complaint, DiningSpace, Reservation, ReservationAction, ReservationOccasion, RestaurantClosure, SiteVisitDaily, StaffProfile, WorkingHours
from .services import AlreadyProcessedError, AvailabilityError, AvailabilityService, ReservationService


class ReservationBaseTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        settings = RestaurantSettings.load()
        settings.minimum_advance_minutes = 0
        settings.maximum_days_ahead = 365
        settings.default_reservation_duration_minutes = 120
        settings.maximum_guests = 30
        settings.save()
        for day in range(7):
            WorkingHours.objects.create(day_of_week=day, open_time=time(10), close_time=time(23))
        cls.space = DiningSpace.objects.create(
            name="VIP-комната", name_uz="VIP xona", space_type=DiningSpace.SpaceType.VIP,
            capacity_min=1, capacity_max=6, is_exclusive=True,
        )
        cls.hall = DiningSpace.objects.create(
            name="Основной зал", capacity_min=1, capacity_max=50, is_exclusive=False,
        )
        user_model = get_user_model()
        cls.manager_user = user_model.objects.create_user("manager", password="safe-test-password", first_name="Анна", is_staff=True)
        cls.manager = StaffProfile.objects.create(user=cls.manager_user, role=StaffProfile.Role.MANAGER, telegram_id=111)
        cls.second_user = user_model.objects.create_user("manager2", password="safe-test-password", is_staff=True)
        cls.second_manager = StaffProfile.objects.create(user=cls.second_user, role=StaffProfile.Role.MANAGER, telegram_id=222)

    def future_date(self, days=2):
        return timezone.localdate() + timedelta(days=days)

    def create_pending(self, *, visit_date=None, visit_time=time(19), guests=6, space=None, **extra):
        with patch("reservations.telegram.notifications.notify_new_reservation"):
            return ReservationService.create(
                customer_name=extra.get("customer_name", "Baxt"), phone=extra.get("phone", "+998 91 777 66 55"),
                date=visit_date or self.future_date(), time=visit_time, guests_count=guests,
                space=space or self.space, occasion=extra.get("occasion", Reservation.Occasion.BIRTHDAY),
                customer_comment=extra.get("customer_comment", "Нужен детский стул"),
            )


class ModelAndAvailabilityTests(ReservationBaseTest):
    def test_dining_space_validates_capacity(self):
        item = DiningSpace(name="Малая", capacity_min=8, capacity_max=4)
        with self.assertRaises(ValidationError):
            item.full_clean()

    def test_staff_profile_roles_and_working_hours(self):
        self.assertEqual(self.manager.get_role_display(), "Administrator")
        self.assertEqual(WorkingHours.objects.count(), 7)
        self.assertIn("Dushanba", str(WorkingHours.objects.get(day_of_week=0)))

    def test_capacity_filter_hides_unsuitable_room(self):
        result = AvailabilityService.spaces_for(self.future_date(), time(19), 10)
        self.assertNotIn(self.space, [space for space, _, _ in result])
        self.assertIn(self.hall, [space for space, _, _ in result])

    def test_closed_day_is_not_bookable(self):
        visit_date = self.future_date()
        RestaurantClosure.objects.create(date=visit_date, reason="Ремонт", full_day=True)
        with self.assertRaisesMessage(AvailabilityError, "В этот день"):
            AvailabilityService.ensure_space_available(self.space, visit_date, time(19), 4)

    def test_working_hours_reject_0100(self):
        with self.assertRaisesMessage(AvailabilityError, "ресторан закрыт"):
            AvailabilityService.ensure_space_available(self.space, self.future_date(), time(1), 4)

    def test_temporary_unavailable_space_has_reason(self):
        self.space.is_temporarily_unavailable = True
        self.space.unavailable_reason = "Ремонт до 25 сентября"
        self.space.save()
        result = AvailabilityService.spaces_for(self.future_date(), time(19), 4)
        row = next(row for row in result if row[0] == self.space)
        self.assertFalse(row[1])
        self.assertEqual(row[2], "Ремонт до 25 сентября")


class ReservationServiceTests(ReservationBaseTest):
    def test_required_example_creates_pending_with_token_and_notification(self):
        mocked_now = timezone.make_aware(datetime(2026, 9, 21, 12, 0))
        with patch("reservations.services.timezone.localtime", return_value=mocked_now), \
             patch("reservations.telegram.notifications.notify_new_reservation") as notify, \
             self.captureOnCommitCallbacks(execute=True):
            reservation = ReservationService.create(
                customer_name="Baxt", phone="+998 90 123 45 67", date=date(2026, 9, 21), time=time(19),
                guests_count=6, space=self.space, occasion=Reservation.Occasion.BIRTHDAY,
                customer_comment="Нужен детский стул",
            )
        self.assertEqual(reservation.status, Reservation.Status.PENDING)
        self.assertTrue(reservation.public_token)
        self.assertRegex(reservation.public_number, r"^R-\d+$")
        notify.assert_called_once_with(reservation.pk)
        self.assertTrue(reservation.actions.filter(action=ReservationAction.Action.CREATED).exists())

    def test_manager_accepts(self):
        reservation = self.create_pending()
        accepted = ReservationService.accept(reservation.pk, self.manager)
        self.assertEqual(accepted.status, Reservation.Status.ACCEPTED)
        self.assertIsNotNone(accepted.accepted_at)
        self.assertEqual(accepted.handled_by_staff, self.manager)

    def test_double_accept_only_first_manager_wins(self):
        reservation = self.create_pending()
        ReservationService.accept(reservation.pk, self.manager)
        with self.assertRaisesMessage(AlreadyProcessedError, "уже обработана"):
            ReservationService.accept(reservation.pk, self.second_manager)
        reservation.refresh_from_db()
        self.assertEqual(reservation.handled_by_staff, self.manager)
        self.assertEqual(reservation.actions.filter(action=ReservationAction.Action.ACCEPTED).count(), 1)

    def test_exclusive_room_overlap_conflicts(self):
        existing = self.create_pending(visit_time=time(18))
        overlapping = self.create_pending(visit_time=time(19))
        ReservationService.accept(existing.pk, self.manager)
        with self.assertRaisesMessage(AvailabilityError, "уже занято"):
            ReservationService.accept(overlapping.pk, self.second_manager)

    def test_exclusive_room_non_conflict_at_2100(self):
        existing = self.create_pending(visit_time=time(18))
        ReservationService.accept(existing.pk, self.manager)
        later = self.create_pending(visit_time=time(21))
        result = ReservationService.accept(later.pk, self.second_manager)
        self.assertEqual(result.status, Reservation.Status.ACCEPTED)

    def test_nonexclusive_hall_allows_parallel_reservations(self):
        first = self.create_pending(space=self.hall, guests=20)
        second = self.create_pending(space=self.hall, guests=20)
        ReservationService.accept(first.pk, self.manager)
        result = ReservationService.accept(second.pk, self.second_manager)
        self.assertEqual(result.status, Reservation.Status.ACCEPTED)

    def test_customer_accepts_proposed_time_after_recheck(self):
        reservation = self.create_pending()
        proposed_date = self.future_date(3)
        ReservationService.propose(reservation.pk, proposed_date=proposed_date, proposed_time=time(20), comment="Можем предложить 20:00", actor=self.manager)
        reservation.refresh_from_db()
        accepted = ReservationService.accept_proposal(reservation)
        self.assertEqual(accepted.status, Reservation.Status.ACCEPTED)
        self.assertEqual(accepted.date, proposed_date)
        self.assertEqual(accepted.time, time(20))

    def test_rejection_reason_persists(self):
        reservation = self.create_pending()
        ReservationService.reject(reservation.pk, reason="Нет свободных мест", actor=self.manager)
        reservation.refresh_from_db()
        self.assertEqual(reservation.status, Reservation.Status.REJECTED)
        self.assertEqual(reservation.rejection_reason, "Нет свободных мест")


class PublicFlowTests(ReservationBaseTest):
    def test_public_page_views_and_unique_visitors_are_counted(self):
        self.client.get(reverse("menu:home"))
        self.client.get(reverse("menu:info"))
        row = SiteVisitDaily.objects.get(date=timezone.localdate())
        self.assertEqual(row.views, 2)
        self.assertEqual(row.unique_visitors, 1)

    def test_reservation_page_and_api_follow_selected_language(self):
        response = self.client.get(reverse("reservations:start"), {"lang": "uz"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Stol band qilish")
        api = self.client.get(reverse("reservations:spaces"), {"date": self.future_date().isoformat(), "time": "19:00", "guests": 4, "lang": "uz"})
        self.assertEqual(api.status_code, 200)
        self.assertTrue(any(row["name"] == "VIP xona" for row in api.json()["spaces"]))
        russian = self.client.get(reverse("reservations:start"), {"lang": "ru"})
        self.assertContains(russian, "Забронировать")

    def test_active_occasions_are_managed_from_database(self):
        option = ReservationOccasion.objects.create(code="ANNIVERSARY", name_ru="Годовщина", name_uz="Yillik", sort_order=15)
        response = self.client.get(reverse("reservations:start"), {"lang": "uz"})
        self.assertContains(response, "Yillik")
        option.is_active = False
        option.save(update_fields=("is_active",))
        response = self.client.get(reverse("reservations:start"), {"lang": "uz"})
        self.assertNotContains(response, "Yillik")

    def test_public_form_creates_reservation(self):
        with patch("reservations.telegram.notifications.notify_new_reservation"), self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(reverse("reservations:start"), {
                "date": self.future_date().isoformat(), "time": "19:00", "guests_count": 4,
                "space": self.space.pk, "customer_name": "Айгуль", "phone": "+998 90 555 44 33",
                "occasion": "", "customer_comment": "У окна",
            })
        reservation = Reservation.objects.get(customer_name="Айгуль")
        self.assertRedirects(response, f"{reverse('reservations:status', args=(reservation.public_token,))}?lang=uz", fetch_redirect_response=False)

    def test_repeated_submission_token_creates_only_one_reservation(self):
        token = uuid.uuid4()
        payload = {
            "date": self.future_date().isoformat(), "time": "19:00", "guests_count": 4,
            "space": self.space.pk, "customer_name": "Double click", "phone": "+998 90 100 20 30",
            "occasion": "", "customer_comment": "", "submission_token": str(token), "lang": "uz",
        }
        with patch("reservations.telegram.notifications.notify_new_reservation"), self.captureOnCommitCallbacks(execute=True):
            first = self.client.post(reverse("reservations:start"), payload)
            second = self.client.post(reverse("reservations:start"), payload)
        self.assertEqual(first.status_code, 302)
        self.assertEqual(second.status_code, 302)
        self.assertEqual(Reservation.objects.filter(submission_token=token).count(), 1)

    def test_status_uses_token_and_never_exposes_phone(self):
        reservation = self.create_pending()
        response = self.client.get(reverse("reservations:status", args=(reservation.public_token,)))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, reservation.phone)
        self.assertEqual(self.client.get(f"/reservation/status/{reservation.pk}/").status_code, 404)
        data = self.client.get(reverse("reservations:status_data", args=(reservation.public_token,))).json()
        self.assertNotIn("phone", data)

    def test_rejected_customer_sees_reason(self):
        reservation = self.create_pending()
        ReservationService.reject(reservation.pk, reason="Выбранное помещение недоступно", actor=self.manager)
        response = self.client.get(reverse("reservations:status", args=(reservation.public_token,)))
        self.assertContains(response, "Выбранное помещение недоступно")


class StaffSecurityTests(ReservationBaseTest):
    def test_normal_account_is_denied_staff_panel(self):
        user = get_user_model().objects.create_user("guest", password="safe-test-password")
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("staff:dashboard")).status_code, 403)

    def test_manager_can_open_and_process_reservation(self):
        reservation = self.create_pending()
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse("staff:reservation_detail", args=(reservation.pk,))).status_code, 200)
        response = self.client.post(reverse("staff:reservation_accept", args=(reservation.pk,)))
        self.assertRedirects(response, reverse("staff:reservation_detail", args=(reservation.pk,)), fetch_redirect_response=False)
        reservation.refresh_from_db()
        self.assertEqual(reservation.status, Reservation.Status.ACCEPTED)
        analytics = self.client.get(reverse("staff:analytics"))
        self.assertEqual(analytics.status_code, 200)
        self.assertContains(analytics, "Просмотры, заявки и гости")

    def test_revoked_manager_bot_access_is_denied(self):
        from telegram_bot.handlers import authorized_profile
        self.assertIsNotNone(async_to_sync(authorized_profile)(111))
        self.manager.is_active = False
        self.manager.save(update_fields=("is_active",))
        self.assertIsNone(async_to_sync(authorized_profile)(111))

    @override_settings(PRIMARY_OWNER_TELEGRAM_ID="999999")
    def test_primary_owner_is_bootstrapped_for_bot(self):
        from telegram_bot.handlers import authorized_profile
        profile = async_to_sync(authorized_profile)(999999)
        self.assertEqual(profile.role, StaffProfile.Role.OWNER)
        self.assertTrue(profile.is_active)
        self.assertTrue(profile.telegram_notifications_enabled)
        self.assertFalse(profile.user.has_usable_password())

    def test_owner_bot_can_add_and_disable_an_administrator(self):
        from telegram_bot.handlers import create_or_restore_admin, deactivate_admin
        profile, success, _ = async_to_sync(create_or_restore_admin)(333333, "Ali")
        self.assertTrue(success)
        self.assertEqual(profile.role, StaffProfile.Role.MANAGER)
        self.assertTrue(profile.is_active)
        async_to_sync(deactivate_admin)(profile.pk)
        profile.refresh_from_db()
        profile.user.refresh_from_db()
        self.assertFalse(profile.is_active)
        self.assertFalse(profile.user.is_active)

    def test_bot_can_toggle_a_room_without_deleting_history(self):
        from telegram_bot.handlers import toggle_space
        space, available = async_to_sync(toggle_space)(self.space.pk, self.manager)
        self.assertFalse(available)
        self.assertTrue(space.is_temporarily_unavailable)
        self.assertTrue(ReservationAction.objects.filter(space=space, action=ReservationAction.Action.SPACE_DISABLED).exists())
        space, available = async_to_sync(toggle_space)(self.space.pk, self.manager)
        self.assertTrue(available)
        self.assertFalse(space.is_temporarily_unavailable)

    def test_manager_cannot_manage_employees(self):
        self.client.force_login(self.manager_user)
        self.assertEqual(self.client.get(reverse("staff:members")).status_code, 403)

    def test_manager_can_add_bilingual_room(self):
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse("staff:space_add"), {
            "name": "Банкетный зал", "name_uz": "Banket zali", "space_type": DiningSpace.SpaceType.HALL,
            "description": "Для праздников", "description_uz": "Bayramlar uchun", "capacity_min": 8,
            "capacity_max": 30, "is_exclusive": "on", "is_active": "on", "is_bookable": "on",
            "sort_order": 30,
        })
        self.assertRedirects(response, reverse("staff:spaces"), fetch_redirect_response=False)
        self.assertTrue(DiningSpace.objects.filter(name_uz="Banket zali", capacity_max=30).exists())

    def test_content_manager_gets_content_only_dashboard(self):
        user = get_user_model().objects.create_user("content", password="safe-test-password", is_staff=True)
        StaffProfile.objects.create(user=user, role=StaffProfile.Role.CONTENT_MANAGER)
        self.client.force_login(user)
        response = self.client.get(reverse("staff:dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Меню и контент")
        self.assertNotContains(response, "Ожидают подтверждения")


class TelegramHistoryTests(ReservationBaseTest):
    def test_history_is_newest_first_and_uses_fifteen_items_per_page(self):
        from telegram_bot.handlers import HISTORY_PAGE_SIZE, _history_page_data

        created = [self.create_pending(customer_name=f"Mehmon {index}") for index in range(17)]
        first_page, total, page, pages = _history_page_data("all", 1)
        second_page, _, second_page_number, _ = _history_page_data("all", 2)

        self.assertEqual(HISTORY_PAGE_SIZE, 15)
        self.assertEqual(total, 17)
        self.assertEqual((page, pages), (1, 2))
        self.assertEqual(second_page_number, 2)
        self.assertEqual(len(first_page), 15)
        self.assertEqual(len(second_page), 2)
        self.assertEqual(first_page[0].pk, created[-1].pk)
        self.assertEqual(second_page[-1].pk, created[0].pk)

    def test_history_filters_keep_rejection_reason(self):
        from telegram_bot.handlers import _history_page_data

        rejected = self.create_pending(customer_name="Rad etilgan mehmon")
        ReservationService.reject(rejected.pk, reason="Bo‘sh joy qolmagan", actor=self.manager)
        items, total, _, _ = _history_page_data("rejected", 1)

        self.assertEqual(total, 1)
        self.assertEqual(items[0].rejection_reason, "Bo‘sh joy qolmagan")

    def test_word_history_report_contains_table_and_status_reason(self):
        from telegram_bot.handlers import HISTORY_LABELS, _history_report_data
        from telegram_bot.reports import build_reservation_report

        rejected = self.create_pending(customer_name="Dilnoza")
        ReservationService.reject(rejected.pk, reason="Tanlangan xona band", actor=self.manager)
        rows, total, truncated = _history_report_data("rejected")
        content = build_reservation_report(
            rows,
            section_title=HISTORY_LABELS["rejected"],
            period_text="Barcha mavjud sanalar",
            generated_at="23.09.2026 18:00",
            total_count=total,
            truncated=truncated,
        )
        document = Document(BytesIO(content))

        self.assertGreater(len(content), 1000)
        self.assertEqual(document.paragraphs[0].text, "Band qilish arizalari hisoboti")
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.tables[0].rows), 2)
        table_text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
        self.assertIn("Rad etildi", table_text)
        self.assertIn("Tanlangan xona band", table_text)
        self.assertIn("Dilnoza", table_text)

    def test_history_date_parser_accepts_two_clear_formats(self):
        from telegram_bot.handlers import parse_history_date

        self.assertEqual(parse_history_date("23.09.2026"), date(2026, 9, 23))
        self.assertEqual(parse_history_date("2026-09-23"), date(2026, 9, 23))
        self.assertIsNone(parse_history_date("ertaga"))


class TelegramComplaintTests(ReservationBaseTest):
    def create_complaint(self, **extra):
        return Complaint.objects.create(
            reason=extra.get("reason", Complaint.Reason.SERVICE),
            space=extra.get("space", self.space),
            place_details=extra.get("place_details", "2-xona"),
            description=extra.get("description", "Ofitsiant uzoq vaqt kelmadi."),
            status=extra.get("status", Complaint.Status.NEW),
        )

    def test_bot_complaints_are_paginated_newest_first(self):
        from telegram_bot.handlers import COMPLAINT_PAGE_SIZE, _complaints_page_data

        created = [self.create_complaint(description=f"Shikoyat {index}") for index in range(17)]
        first_page, total, page, pages = _complaints_page_data("new", 1)
        second_page, _, _, _ = _complaints_page_data("new", 2)
        self.assertEqual(COMPLAINT_PAGE_SIZE, 15)
        self.assertEqual((total, page, pages), (17, 1, 2))
        self.assertEqual(first_page[0].pk, created[-1].pk)
        self.assertEqual(second_page[-1].pk, created[0].pk)

    def test_bot_can_review_and_resolve_complaint(self):
        from telegram_bot.handlers import update_complaint_status

        complaint = self.create_complaint()
        async_to_sync(update_complaint_status)(complaint.pk, Complaint.Status.IN_REVIEW, self.manager)
        complaint.refresh_from_db()
        self.assertEqual(complaint.status, Complaint.Status.IN_REVIEW)
        self.assertEqual(complaint.handled_by_staff, self.manager)
        self.assertIsNotNone(complaint.reviewed_at)

        async_to_sync(update_complaint_status)(complaint.pk, Complaint.Status.RESOLVED, self.manager)
        complaint.refresh_from_db()
        self.assertEqual(complaint.status, Complaint.Status.RESOLVED)
        self.assertIsNotNone(complaint.resolved_at)
